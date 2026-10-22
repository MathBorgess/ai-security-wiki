#!/usr/bin/env python3
"""Extract public reference tabs only. Never import assessment/template fields.

Usage: python scripts/import_compendium.py SOURCE.xlsx --date YYYY-MM-DD
The source must be an export of the official public workbook, not a user assessment.
No downloads, macro execution, or formula execution occur here.
"""
import argparse
import datetime
import hashlib
import json
from pathlib import Path
import re
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SOURCE_ID = "1lD5nKC1QNoUTDXdMp31NA9eKcUYGjhtC34MCfQpqpGo"
TABS = [
    (0, "Data and AI Security Framework Introduction", 785337750, 28, 1),
    (1, "Table of Content", 267581052, 119, 8),
    (2, "AI Lifecycle Risks", 307775890, 99, 31),
    (3, "Databricks AI Mitigation Controls", 780486953, 75, 28),
    (6, "Third-Party Tools", 1670825065, 51, 3),
]


def clean(value):
    if value is None:
        return ""
    if not isinstance(value, (str, int, float, bool)):
        raise ValueError(f"Unsupported cell value type: {type(value).__name__}")
    text = str(value)
    # XLSX escapes XML-forbidden control characters. Normalize text only;
    # do not reinterpret formulas, HTML, or instructions from source cells.
    text = re.sub(r"_x([0-9a-fA-F]{4})_", lambda m: chr(int(m[1], 16)), text)
    return text.replace("\r\n", "\n").replace("\r", "\n").replace("\x0b", "\n")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("--date", required=True, type=datetime.date.fromisoformat)
    args = parser.parse_args()
    raw = args.xlsx.read_bytes()
    if len(raw) > 10_000_000:
        raise ValueError("Expected a compendium export smaller than 10 MB")
    workbook = openpyxl.load_workbook(args.xlsx, data_only=True, keep_links=False)
    if len(workbook.worksheets) != 7:
        raise ValueError("Worksheet layout changed; review the importer first")
    sheets = []
    for index, title, gid, max_row, max_col in TABS:
        sheet = workbook.worksheets[index]
        if not title.startswith(sheet.title.strip()):
            raise ValueError(f"Unexpected tab {index}: {sheet.title}")
        if any(c.value for row in sheet.iter_rows(min_row=max_row + 1) for c in row):
            raise ValueError(f"New rows in {title}; review counts before import")
        if sheet.max_column > max_col and any(
            c.value for row in sheet.iter_rows(min_col=max_col + 1) for c in row
        ):
            raise ValueError(f"New columns in {title}; review schema before import")
        rows, hyperlinks, notes = [], {}, {}
        for row in sheet.iter_rows(max_row=max_row, max_col=max_col):
            rows.append([clean(c.value) for c in row])
            for cell in row:
                if cell.hyperlink and cell.hyperlink.target:
                    url = cell.hyperlink.target
                    if not url.startswith(("https://", "http://")):
                        raise ValueError(f"Unsafe hyperlink in {title}:{cell.coordinate}")
                    hyperlinks[cell.coordinate] = url
                if cell.comment:
                    notes[cell.coordinate] = clean(cell.comment.text)
        sheets.append(dict(title=title, gid=gid, rows=rows, hyperlinks=hyperlinks, notes=notes))
    if sheets[2]["rows"][0][0] != "Risk ID" or sheets[3]["rows"][0][0] != "Control ID":
        raise ValueError("Headers changed")
    if sum(bool(r[0]) for r in sheets[2]["rows"][2:]) != 97:
        raise ValueError("Expected 97 risks")
    if sum(bool(r[0]) for r in sheets[3]["rows"][2:]) != 73:
        raise ValueError("Expected 73 controls")
    result = dict(
        schema_version=1, framework="DASF", version="3.0",
        retrieved_on=str(args.date), source_url=f"https://docs.google.com/spreadsheets/d/{SOURCE_ID}/edit",
        source_xlsx_sha256=hashlib.sha256(raw).hexdigest(),
        exclusions=["DASF Risk Applicability", "DASF Control Applicability"],
        exclusion_reason="These are assessment templates; applicability is derived from the seven reference columns.",
        sheets=sheets,
    )
    destination = ROOT / "sources/compendium.json"
    destination.parent.mkdir(exist_ok=True)
    destination.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    print(f"Extracted 97 risks, 73 controls and public reference tabs to {destination}")


if __name__ == "__main__":
    main()
